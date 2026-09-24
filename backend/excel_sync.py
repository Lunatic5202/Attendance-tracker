"""Attendance -> Excel workbook sync.

Builds a real ``.xlsx`` workbook with ``openpyxl`` and pushes it to the owner's
OneDrive for Business via Microsoft Graph. The workbook is fully regenerated on
every sync (single source of truth = SQLite), so it is always correct and never
depends on state that a Render redeploy might wipe.

Schema written by: creators / public-facing wrappers only. Attendance timestamps
and employee profile fields are decrypted here for the export (they are still
encrypted at rest in SQLite). Face templates are NEVER exported.
"""

from __future__ import annotations

import io
import logging
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from backend import onedrive
from backend.security import public_attendance, public_employee
from backend.store import store

log = logging.getLogger("excel_sync")

HEADER_FILL = PatternFill(start_color="1F3864", end_color="1F3864", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)
_TITLE_FONT = Font(size=14, bold=True)

ATTENDANCE_COLUMNS = [
    ("date", "Date", 12),
    ("employee_id", "Employee ID", 12),
    ("name", "Name", 26),
    ("department", "Department", 20),
    ("check_in", "Check-In", 22),
    ("check_out", "Check-Out", 22),
    ("hours", "Hours Worked", 14),
    ("status", "Status", 11),
    ("source", "Source", 9),
]

EMPLOYEE_COLUMNS = [
    ("id", "Employee ID", 12),
    ("name", "Name", 26),
    ("department", "Department", 20),
    ("role", "Role", 20),
    ("email", "Email", 30),
    ("phone", "Phone", 16),
    ("face_enrolled", "Face Enrolled", 14),
    ("is_active", "Active", 10),
]


def _fmt_time(value: str | None) -> str:
    if not value:
        return ""
    if len(value) >= 5:
        return value[:5]
    return value


def _add_sheet(wb: Workbook, title: str, columns, rows) -> int:
    ws = wb.create_sheet(title=title)
    for col, (_, header, width) in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.row_dimensions[1].height = 20
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}1"
    for r, record in enumerate(rows, start=2):
        for col, (key, _, _) in enumerate(columns, start=1):
            ws.cell(row=r, column=col, value=record.get(key, ""))
    return r - 1 if rows else 0


def build_workbook() -> tuple[bytes, dict]:
    """Return ``(xlsx_bytes, info)`` for the current attendance snapshot."""
    attendance_rows = [public_attendance(row) for row in store.attendance_all()]
    for row in attendance_rows:
        row["check_in"] = _fmt_time(row.get("check_in"))
        row["check_out"] = _fmt_time(row.get("check_out"))

    employee_rows = [public_employee(row) for row in store.employees()]

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"

    built_since = "Every scan since the system was deployed"
    ws["A1"] = "Attendance Tracker"
    ws["A1"].font = _TITLE_FONT
    ws["A2"] = built_since
    ws["A4"] = "Sheet"
    ws["B4"] = "Rows"
    for cell in (ws["A4"], ws["B4"]):
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
    ws["A5"] = "Attendance"
    ws["B5"] = len(attendance_rows)
    ws["A6"] = "Employees"
    ws["B6"] = len(employee_rows)
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 12

    attendance_count = _add_sheet(wb, "Attendance", ATTENDANCE_COLUMNS, attendance_rows)
    employee_count = _add_sheet(wb, "Employees", EMPLOYEE_COLUMNS, employee_rows)

    buffer = io.BytesIO()
    wb.save(buffer)
    info = {
        "attendance_rows": attendance_count,
        "employee_rows": employee_count,
        "sheets": ["Summary", "Attendance", "Employees"],
    }
    return buffer.getvalue(), info


def sync_to_excel() -> dict:
    """Regenerate the workbook and upload it to the configured OneDrive folder."""
    if not onedrive.configured():
        return {
            "ok": False,
            "disabled": True,
            "reason": "Microsoft Graph is not configured (set MS_CLIENT_ID / MS_CLIENT_SECRET / "
                      "MS_TENANT_ID / MS_DRIVE_UPN).",
        }
    filename = os.getenv("MS_EXCEL_FILENAME", "Attendance.xlsx")
    folder = os.getenv("MS_FOLDER", "Attendance Tracker")
    data, info = build_workbook()
    uploaded = onedrive.upload(
        folder,
        filename,
        data,
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    info.update({"ok": True, "filename": uploaded["name"], "url": uploaded.get("url"),
                 "folder": folder})
    return info